#!/usr/bin/env python3
"""Rellena el Excel oficial de casos de prueba de Idealista (Contact/Property/
Images) con los resultados reales de scripts/idealista-run-official-testcases.mts.

Uso:
    python3 scripts/idealista-fill-testcases-xlsx.py <original.xlsx> <results.json> <salida.xlsx>

Sólo escribe en las columnas "Actual Result" / propertyId-contactId / "JSON sent"
/ "Comments" de cada fila cuyo ID (columna A) coincide con un resultado. No toca
nada más del fichero (columnas, formato, encabezados de Idealista).
"""
import json
import sys

import openpyxl


def normalize(test_id: str) -> str:
    return test_id.replace(" ", "").replace("\n", "").strip().lower()


def main() -> None:
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(1)

    original_path, results_path, out_path = sys.argv[1:4]

    with open(results_path, encoding="utf-8") as f:
        results = json.load(f)

    by_id = {}
    for r in results:
        by_id.setdefault(normalize(r["id"]), []).append(r)

    wb = openpyxl.load_workbook(original_path)

    id_col = {"Contact": "F", "Property": "F", "Images": None}  # F = contactId/propertyId; Images no tiene esa columna
    actual_col = "E"
    json_col = {"Contact": "G", "Property": "G", "Images": "F"}
    comment_col = {"Contact": "H", "Property": "H", "Images": "G"}

    filled = 0
    unmatched_rows = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
            test_id_cell = row[0]  # columna A
            if not test_id_cell.value:
                continue
            key = normalize(str(test_id_cell.value))
            matches = by_id.get(key)
            if not matches:
                unmatched_rows.append((sheet_name, test_id_cell.value))
                continue

            r = matches[0]  # una llamada por caso de prueba
            row_num = test_id_cell.row

            actual_text = f"HTTP {r['status']}"
            if r.get("message"):
                actual_text += f" — {r['message']}"
            ws[f"{actual_col}{row_num}"] = actual_text

            id_column = id_col.get(sheet_name)
            if id_column and r.get("entityId") is not None:
                ws[f"{id_column}{row_num}"] = r["entityId"]

            body = r.get("requestBody")
            if body is not None:
                ws[f"{json_col[sheet_name]}{row_num}"] = json.dumps(body, ensure_ascii=False, indent=2)

            comment = r.get("comment") or ""
            ws[f"{comment_col[sheet_name]}{row_num}"] = comment

            filled += 1

    wb.save(out_path)
    print(f"Filas rellenadas: {filled}")
    if unmatched_rows:
        print(f"Filas sin resultado (no ejecutadas): {len(unmatched_rows)}")
        for sheet, tid in unmatched_rows:
            print(f"  - [{sheet}] {tid!r}")

    used_ids = {normalize(r["id"]) for r in results}
    all_row_ids = set()
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
            if row[0].value:
                all_row_ids.add(normalize(str(row[0].value)))
    extra = used_ids - all_row_ids
    if extra:
        print(f"AVISO: resultados sin fila correspondiente en el xlsx: {extra}")


if __name__ == "__main__":
    main()
